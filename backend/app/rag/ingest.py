"""Pipeline de ingesta documental: archivo → texto → chunks → embeddings → pgvector."""

from __future__ import annotations

import io
import re
from dataclasses import dataclass

import yaml
from sqlalchemy.orm import Session

from app.models import RagChunk, RagDocument
from app.rag.embedder import embed

_CHUNK_SIZE = 1100


@dataclass(frozen=True)
class IngestResult:
    doc_id: int
    titulo: str
    chunks: int


def _extract_text_pdf(content: bytes) -> str:
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(content))
    return "\n\n".join(page.extract_text() or "" for page in reader.pages)


def _extract_text_docx(content: bytes) -> str:
    from docx import Document
    doc = Document(io.BytesIO(content))
    parts: list[str] = []
    for block in doc.element.body:
        tag = block.tag.split("}")[-1]
        if tag == "p":
            text = block.text_content().strip()
            if text:
                parts.append(text)
        elif tag == "tbl":
            from docx.oxml.ns import qn
            rows = block.findall(f".//{qn('w:tr')}")
            for row in rows:
                cells = [c.text_content().strip() for c in row.findall(f".//{qn('w:tc')}")]
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _extract_text_excel(content: bytes, filename: str) -> list[str]:
    """Chunks por grupos de filas. Repite la fila de cabecera en cada chunk."""
    import openpyxl
    ROWS_PER_CHUNK = 12
    chunks: list[str] = []
    wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    for sheet in wb.worksheets:
        rows: list[str] = []
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
            if cells:
                rows.append(" | ".join(cells))
        if not rows:
            continue
        header_row = rows[0]
        body = rows[1:] if len(rows) > 1 else rows
        prefix = f"[Hoja: {sheet.title}]\n{header_row}"
        for i in range(0, len(body), ROWS_PER_CHUNK):
            group = body[i : i + ROWS_PER_CHUNK]
            chunks.append(prefix + "\n" + "\n".join(group))
    return chunks


_FRONTMATTER_RE = re.compile(r"^---\n(.*?)\n---\n", re.DOTALL)
_TAG_LINE_RE = re.compile(r"^\s*#[\wÀ-ÿ/_-]+(\s+#[\wÀ-ÿ/_-]+)*\s*$")
_WIKILINK_RE = re.compile(r"\[\[([^\]|]+)(?:\|([^\]]+))?\]\]")
_CALLOUT_HEADER_RE = re.compile(r"^\s*>\s*\[![a-zA-Z]+\]")


def clean_markdown(raw: str) -> tuple[dict, str]:
    """Devuelve (frontmatter, cuerpo limpio) de una nota Obsidian.

    Elimina ruido que no debe embeberse: frontmatter YAML, líneas de etiquetas
    (#tag), el callout de aviso "Contenido público", la sección "Notas relacionadas"
    (navegación interna) y convierte [[enlaces|alias]] en su texto visible.
    """
    meta: dict = {}
    fm = _FRONTMATTER_RE.match(raw)
    body = raw
    if fm:
        try:
            meta = yaml.safe_load(fm.group(1)) or {}
        except yaml.YAMLError:
            meta = {}
        body = raw[fm.end():]

    lines = body.splitlines()
    out: list[str] = []
    skip_callout = False
    skip_section = False
    for line in lines:
        stripped = line.strip()

        # Corta una sección de navegación interna hasta el siguiente encabezado.
        if re.match(r"^#{1,6}\s+Notas relacionadas\b", stripped, re.IGNORECASE):
            skip_section = True
            continue
        if skip_section:
            if stripped.startswith("#") and re.match(r"^#{1,6}\s+", stripped):
                skip_section = False  # nuevo encabezado: deja de saltar
            else:
                continue

        # Elimina el callout de aviso a editores ("Contenido público").
        if _CALLOUT_HEADER_RE.match(line):
            skip_callout = "público" in line.lower() or "interno" in line.lower()
            if skip_callout:
                continue
        if skip_callout:
            if stripped.startswith(">"):
                continue
            skip_callout = False

        # Quita líneas que son solo etiquetas (#empresa #faq), no encabezados.
        if _TAG_LINE_RE.match(line) and not re.match(r"^#{1,6}\s+\S", stripped):
            continue

        # [[enlace|alias]] -> alias ; [[enlace]] -> enlace
        line = _WIKILINK_RE.sub(lambda m: m.group(2) or m.group(1), line)
        # Desenvuelve blockquotes que no eran callouts de aviso.
        line = re.sub(r"^\s*>\s?", "", line)
        out.append(line)

    cleaned = "\n".join(out)
    cleaned = re.sub(r"\n{3,}", "\n\n", cleaned).strip()
    return meta, cleaned


def _extract_text(content: bytes, filename: str) -> str:
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext == "pdf":
        return _extract_text_pdf(content)
    if ext == "docx":
        return _extract_text_docx(content)
    if ext in {"md", "markdown"}:
        return clean_markdown(content.decode("utf-8", errors="replace"))[1]
    return content.decode("utf-8", errors="replace")


def _chunk_text(text: str) -> list[str]:
    """Chunking por párrafos hasta _CHUNK_SIZE con solapamiento de un párrafo."""
    paragraphs = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
    if not paragraphs:
        stripped = text.strip()
        return [stripped] if stripped else []

    chunks: list[str] = []
    current: list[str] = []
    length = 0

    for p in paragraphs:
        if len(p) > _CHUNK_SIZE:
            if current:
                chunks.append("\n\n".join(current))
                current, length = [], 0
            for i in range(0, len(p), _CHUNK_SIZE):
                piece = p[i : i + _CHUNK_SIZE].strip()
                if piece:
                    chunks.append(piece)
            continue

        if current and length + len(p) + 2 > _CHUNK_SIZE:
            chunks.append("\n\n".join(current))
            current = current[-1:]
            length = len(current[0]) + 2 if current else 0

        current.append(p)
        length += len(p) + 2

    if current:
        chunks.append("\n\n".join(current))
    return chunks


def _store_chunks(
    session: Session,
    chunks: list[str],
    titulo: str,
    fuente: str,
    tipo: str | None,
    area: str | None,
    content_hash: str | None,
    nivel_confidencialidad: str = "interno",
) -> IngestResult:
    if not chunks:
        raise ValueError("El documento no contiene texto extraíble.")

    # Generar los embeddings ANTES de tocar la DB: si la llamada al proveedor falla,
    # no debe quedar un documento a medias (sin chunks) que luego se comitee.
    embeddings = embed(chunks)

    doc = RagDocument(
        titulo=titulo,
        tipo=tipo,
        area=area,
        fuente_original=fuente,
        content_hash=content_hash,
        nivel_confidencialidad=nivel_confidencialidad,
    )
    session.add(doc)
    session.flush()

    for i, (texto, emb) in enumerate(zip(chunks, embeddings)):
        session.add(RagChunk(doc_id=doc.id, texto=texto, pagina=i, embedding=emb))

    session.commit()
    return IngestResult(doc_id=doc.id, titulo=titulo, chunks=len(chunks))


def ingest_document(
    session: Session,
    content: bytes,
    filename: str,
    titulo: str,
    tipo: str | None = None,
    area: str | None = None,
    content_hash: str | None = None,
) -> IngestResult:
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext in {"xlsx", "xls"}:
        chunks = _extract_text_excel(content, filename)
    else:
        chunks = _chunk_text(_extract_text(content, filename))
    return _store_chunks(session, chunks, titulo, filename, tipo, area, content_hash)


def ingest_text(
    session: Session,
    text: str,
    titulo: str,
    fuente: str,
    tipo: str | None = None,
    area: str | None = None,
    content_hash: str | None = None,
    nivel_confidencialidad: str = "interno",
) -> IngestResult:
    """Indexa texto ya extraído y limpio (p. ej. una nota markdown)."""
    chunks = _chunk_text(text)
    return _store_chunks(
        session, chunks, titulo, fuente, tipo, area, content_hash, nivel_confidencialidad
    )
